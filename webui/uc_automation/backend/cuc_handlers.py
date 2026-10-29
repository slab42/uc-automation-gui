"""CUC call handlers: classify, menu entries, CSV in/out, call tree workbook.

Written for this app (cuc/list_CallHandlers.py, list_NonSubscriber_CallHandlers_WithMenuEntries.py,
build_call_trees.py and delete_CallHandlers.py were used as reference only). The CSV layouts match those
scripts, so files move between the CLI and the web app.
"""

import csv
import io
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor

from .cuc_api import CucClient, CucError

SKIP_HANDLERS = {"Opening Greeting", "Operator", "Goodbye", "undeliverablemessagesmailbox", "operator"}
MAX_WORKERS = 10
MAX_DEPTH = 50

HANDLER_FIELDS = ["displayName", "extension", "objectId", "uri"]
MENU_FIELDS = [
    "sourceHandler", "sourceExtension", "sourceObjectId", "touchToneKey", "locked", "actionCode",
    "destinationType", "destinationName", "targetConversation", "targetHandlerObjectId", "targetHandlerName",
    "transferDisplayName", "transferNumber", "transferType", "transferRings", "menuEntryObjectId", "menuEntryUri",
]
ERROR_FIELDS = ["sourceHandler", "sourceExtension", "sourceObjectId", "menuUrl", "error"]
_HANDLER_REQUIRED = {"displayName", "extension", "objectId", "uri"}
_MENU_REQUIRED = {"sourceHandler", "sourceExtension", "sourceObjectId", "touchToneKey", "destinationType",
                  "destinationName", "targetHandlerObjectId"}


def _clean(value) -> str:
    return str(value).strip() if value is not None else ""


def to_csv(rows: list[dict], fields: list[str]) -> str:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=fields, extrasaction="ignore", lineterminator="\n")
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue()


def read_csv(data: bytes) -> list[dict[str, str]]:
    text = data.decode("utf-8-sig")
    return [{str(k).strip(): (v or "").strip() for k, v in row.items() if k is not None}
            for row in csv.DictReader(io.StringIO(text))]


def _require(rows: list[dict[str, str]], required: set[str], label: str) -> None:
    if not rows:
        raise ValueError(f"The {label} CSV has no data rows.")
    missing = required.difference(rows[0])
    if missing:
        raise ValueError(f"The {label} CSV is missing columns: {', '.join(sorted(missing))}")


def parse_handler_csv(data: bytes) -> list[dict[str, str]]:
    """Call handler rows for Map (needs displayName, extension, objectId, uri)."""
    rows = read_csv(data)
    _require(rows, _HANDLER_REQUIRED, "call handler")
    return rows


def parse_menu_csv(data: bytes) -> list[dict[str, str]]:
    rows = read_csv(data)
    _require(rows, _MENU_REQUIRED, "menu entry")
    return rows


def parse_delete_csv(data: bytes) -> list[dict[str, str]]:
    """Rows to delete: needs objectId; displayName and extension are optional."""
    rows = [r for r in read_csv(data) if r.get("objectId")]
    return [{"displayName": r.get("displayName", ""), "extension": r.get("extension", ""),
             "objectId": r["objectId"]} for r in rows]


# --- listing ---

def classify(items: list[dict[str, str]], search: str = "") -> tuple[list[dict[str, str]], dict[str, dict]]:
    """(non-subscriber handlers to export, every handler by objectId) from raw call handler dicts.

    Subscriber (user) handlers and the built-in SKIP_HANDLERS are left out of the first list.
    `search` keeps only display names containing the text, case-insensitive.
    """
    needle = search.strip().casefold()
    handlers, by_id = [], {}
    for item in items:
        oid = _clean(item.get("ObjectId"))
        name = _clean(item.get("DisplayName"))
        is_sub = bool(item.get("RecipientSubscriberObjectId"))
        if oid:
            by_id[oid] = {"displayName": name, "isSubscriber": is_sub}
        if is_sub or name in SKIP_HANDLERS or (needle and needle not in name.casefold()):
            continue
        handlers.append({"displayName": name, "extension": _clean(item.get("DtmfAccessId")), "objectId": oid,
                         "uri": _clean(item.get("URI"))})
    handlers.sort(key=lambda r: (r["displayName"].casefold(), r["objectId"]))
    return handlers, by_id


def list_handlers(client: CucClient, search: str = "") -> tuple[list[dict[str, str]], dict[str, dict]]:
    return classify(client.list_call_handlers(), search)


def _destination(entry: dict, target_name: str, subscriber_name: str) -> tuple[str, str]:
    transfer = _clean(entry.get("TransferNumber"))
    target_id = _clean(entry.get("TargetHandlerObjectId"))
    conversation = _clean(entry.get("TargetConversation"))
    if transfer:
        return "Transfer", _clean(entry.get("DisplayName")) or transfer
    if target_id:
        if target_name:
            return "Call Handler", target_name
        if subscriber_name:
            return "Subscriber Mailbox", subscriber_name
        return "Unknown Handler", target_id
    if conversation:
        return "Conversation", conversation
    return "No configured destination", ""


def collect_menu_entries(client: CucClient, handlers: list[dict[str, str]],
                         all_by_id: dict[str, dict]) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """(menu entry rows, error rows). Entries are fetched in parallel, one request per handler."""
    name_by_id = {h["objectId"]: h["displayName"] for h in handlers if h["objectId"]}

    def fetch(h):
        try:
            return client.menu_entries(h["objectId"]), ""
        except CucError as e:
            return [], str(e)
        except Exception as e:
            return [], f"Unexpected error: {e}"

    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as pool:
        fetched = list(pool.map(fetch, handlers))

    rows, errors = [], []
    for h, (entries, err) in zip(handlers, fetched):
        if err:
            errors.append({"sourceHandler": h["displayName"], "sourceExtension": h["extension"],
                           "sourceObjectId": h["objectId"],
                           "menuUrl": f"handlers/callhandlers/{h['objectId']}/menuentries", "error": err})
            continue
        for e in entries:
            target_id = _clean(e.get("TargetHandlerObjectId"))
            target_name = name_by_id.get(target_id, "")
            sub_name = ""
            if not target_name and all_by_id.get(target_id, {}).get("isSubscriber"):
                sub_name = all_by_id[target_id]["displayName"]
            dtype, dname = _destination(e, target_name, sub_name)
            locked = _clean(e.get("Locked")).lower()
            rows.append({
                "sourceHandler": h["displayName"], "sourceExtension": h["extension"], "sourceObjectId": h["objectId"],
                "touchToneKey": _clean(e.get("TouchtoneKey")),
                "locked": {"true": "TRUE", "false": "FALSE"}.get(locked, ""),
                "actionCode": _clean(e.get("Action")), "destinationType": dtype, "destinationName": dname,
                "targetConversation": _clean(e.get("TargetConversation")), "targetHandlerObjectId": target_id,
                "targetHandlerName": target_name, "transferDisplayName": _clean(e.get("DisplayName")),
                "transferNumber": _clean(e.get("TransferNumber")), "transferType": _clean(e.get("TransferType")),
                "transferRings": _clean(e.get("TransferRings")), "menuEntryObjectId": _clean(e.get("ObjectId")),
                "menuEntryUri": _clean(e.get("URI")),
            })
    order = {str(i): i for i in range(10)}
    order.update({"*": 10, "#": 11})
    rows.sort(key=lambda r: (r["sourceHandler"].casefold(), order.get(r["touchToneKey"], 99),
                             r["destinationName"].casefold()))
    return rows, errors


# --- call tree workbook ---

def _key_rank(value: str) -> tuple[int, str]:
    order = {str(i): i for i in range(10)}
    order.update({"*": 10, "#": 11})
    return order.get(value, 99), value


def _is_routable(row: dict[str, str]) -> bool:
    return bool(row.get("targetHandlerObjectId") or row.get("transferNumber") or row.get("targetConversation")
                or (row.get("destinationType") and row.get("destinationType") != "No configured destination"))


def _tree_rows(root_id, handlers_by_id, edges_by_source, issues) -> list[dict]:
    root = handlers_by_id[root_id]
    root_name = root.get("displayName", root_id)
    rows: list[dict] = []

    def walk(current_id, level, path_ids, path_text):
        current_name = handlers_by_id.get(current_id, {}).get("displayName", current_id)
        if level > MAX_DEPTH:
            issues.append({"issueType": "Maximum depth exceeded", "sourceHandler": current_name,
                           "sourceObjectId": current_id, "touchToneKey": "", "targetHandler": "", "targetObjectId": "",
                           "details": f"Tree traversal exceeded MAX_DEPTH={MAX_DEPTH}"})
            return
        for edge in sorted(edges_by_source.get(current_id, []), key=lambda r: _key_rank(r.get("touchToneKey", ""))):
            key = edge.get("touchToneKey", "")
            dtype = edge.get("destinationType", "")
            target_id = edge.get("targetHandlerObjectId", "")
            dest = (edge.get("targetHandlerName") or edge.get("destinationName") or edge.get("transferDisplayName")
                    or edge.get("transferNumber") or edge.get("targetConversation") or target_id)
            branch = f"Key {key} -> {dest}" if key else dest
            child_path = f"{path_text} -> {branch}"
            status = ""
            if target_id and target_id in path_ids:
                status = "CIRCULAR REFERENCE"
            elif target_id and target_id not in handlers_by_id and dtype != "Subscriber Mailbox":
                status = "TARGET NOT FOUND"
            rows.append({
                "rootHandler": root_name, "rootExtension": root.get("extension", ""), "level": level,
                "tree": f"{'    ' * level}|- {branch}", "path": child_path, "sourceHandler": current_name,
                "sourceObjectId": current_id, "touchToneKey": key, "destinationType": dtype,
                "destinationName": dest, "transferNumber": edge.get("transferNumber", ""),
                "targetHandlerObjectId": target_id, "status": status,
            })
            if status:
                issues.append({"issueType": status.title(), "sourceHandler": current_name,
                               "sourceObjectId": current_id, "touchToneKey": key, "targetHandler": dest,
                               "targetObjectId": target_id, "details": child_path})
                continue
            if target_id and dtype == "Call Handler":
                walk(target_id, level + 1, path_ids + (target_id,), child_path)

    rows.append({
        "rootHandler": root_name, "rootExtension": root.get("extension", ""), "level": 0, "tree": root_name,
        "path": root_name, "sourceHandler": root_name, "sourceObjectId": root_id, "touchToneKey": "",
        "destinationType": "Root Call Handler", "destinationName": root_name, "transferNumber": "",
        "targetHandlerObjectId": "", "status": "",
    })
    walk(root_id, 1, (root_id,), root_name)
    return rows


def build_call_trees(handlers: list[dict[str, str]], menu_entries: list[dict[str, str]]) -> tuple[bytes, dict[str, int]]:
    """The call tree workbook (xlsx bytes) and its summary counts."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    caution = PatternFill("solid", fgColor="FCE4D6")
    error = PatternFill("solid", fgColor="F4CCCC")

    handlers_by_id = {r["objectId"]: r for r in handlers if r.get("objectId")}
    routable = [r for r in menu_entries if _is_routable(r)]
    edges: dict[str, list] = defaultdict(list)
    targeted: set[str] = set()
    issues: list[dict] = []
    for r in routable:
        src, tgt = r.get("sourceObjectId", ""), r.get("targetHandlerObjectId", "")
        if src:
            edges[src].append(r)
        if tgt:
            targeted.add(tgt)
            if tgt not in handlers_by_id:
                issues.append({"issueType": "Target not found", "sourceHandler": r.get("sourceHandler", ""),
                               "sourceObjectId": src, "touchToneKey": r.get("touchToneKey", ""),
                               "targetHandler": r.get("targetHandlerName", "") or r.get("destinationName", ""),
                               "targetObjectId": tgt,
                               "details": "Target ObjectId is not present in the Call Handlers list"})
    sources = set(edges)
    root_ids = sorted(sources - targeted, key=lambda o: handlers_by_id.get(o, {}).get("displayName", o).casefold())

    covered: set[str] = set()
    tree: list[dict] = []
    for rid in root_ids:
        part = _tree_rows(rid, handlers_by_id, edges, issues) if rid in handlers_by_id else []
        tree.extend(part)
        covered.update(str(r["sourceObjectId"]) for r in part if r.get("sourceObjectId"))
    for sid in sorted(sources - covered):
        if sid in handlers_by_id:  # fully cyclic component with no natural root
            tree.extend(_tree_rows(sid, handlers_by_id, edges, issues))

    roots = [{"displayName": handlers_by_id[o].get("displayName", ""), "extension": handlers_by_id[o].get("extension", ""),
              "objectId": o, "outboundRoutes": len(edges.get(o, []))} for o in root_ids if o in handlers_by_id]
    orphans = [{"displayName": r.get("displayName", ""), "extension": r.get("extension", ""), "objectId": o,
                "reason": "Not a source and not targeted by another exported call handler"}
               for o, r in handlers_by_id.items() if o not in sources and o not in targeted]
    orphans.sort(key=lambda r: r["displayName"].casefold())

    wb = Workbook()
    wb.remove(wb.active)
    summary = wb.create_sheet("Summary")
    for row in [("Cisco Unity Call Tree Workbook", ""), ("Call handlers imported", len(handlers)),
                ("Menu-entry rows imported", len(menu_entries)), ("Routable menu entries", len(routable)),
                ("Root handlers identified", len(root_ids)), ("Orphan handlers", len(orphans)),
                ("Issues detected", len(issues)), ("Maximum traversal depth", MAX_DEPTH)]:
        summary.append(row)
    summary["A1"].fill = PatternFill("solid", fgColor="D9EAF7")
    summary["A1"].font = Font(bold=True, size=14)
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 18
    summary.sheet_view.showGridLines = False

    def sheet(name, table_name, headers, rows):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for r in rows:
            ws.append([r.get(h, "") for h in headers])
        for cell in ws[1]:
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        if rows:
            t = Table(displayName=table_name, ref=ws.dimensions)
            t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(t)
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 60)
        return ws

    sheet("Call Handlers", "CallHandlersTable", ["displayName", "extension", "objectId", "uri"], handlers)
    sheet("Menu Entries", "MenuEntriesTable", list(menu_entries[0]) if menu_entries else MENU_FIELDS, menu_entries)
    sheet("Root Handlers", "RootHandlersTable", ["displayName", "extension", "objectId", "outboundRoutes"], roots)
    tree_headers = ["rootHandler", "rootExtension", "level", "tree", "path", "sourceHandler", "sourceObjectId",
                    "touchToneKey", "destinationType", "destinationName", "transferNumber",
                    "targetHandlerObjectId", "status"]
    ws = sheet("Call Trees", "CallTreesTable", tree_headers, tree)
    ws.column_dimensions["D"].width = 70
    ws.column_dimensions["E"].width = 100
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 4).alignment = Alignment(indent=min(int(ws.cell(row, 3).value or 0), 15))
        if ws.cell(row, 13).value:
            for col in range(1, 14):
                ws.cell(row, col).fill = error
    sheet("Orphan Handlers", "OrphanHandlersTable", ["displayName", "extension", "objectId", "reason"], orphans)
    issue_headers = ["issueType", "sourceHandler", "sourceObjectId", "touchToneKey", "targetHandler",
                     "targetObjectId", "details"]
    ws = sheet("Issues", "IssuesTable", issue_headers, issues)
    for row in range(2, ws.max_row + 1):
        for col in range(1, len(issue_headers) + 1):
            ws.cell(row, col).fill = caution

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), {
        "handlers": len(handlers), "menuEntries": len(menu_entries), "routable": len(routable),
        "roots": len(root_ids), "orphans": len(orphans), "issues": len(issues), "treeRows": len(tree),
    }
